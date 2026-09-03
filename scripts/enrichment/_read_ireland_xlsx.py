"""Read Ireland charities XLSX."""
import openpyxl
import csv

wb = openpyxl.load_workbook(r'E:\grid\data\ireland_charities_register.xlsx', read_only=True, data_only=True)
print('Sheets:', wb.sheetnames)

ws = wb['Public Register']
print(f'Sheet: Public Register')

# Headers from first row
headers = []
for cell in next(ws.iter_rows(min_row=1, max_row=1)):
    headers.append(cell.value)
print(f'\nColumns ({len(headers)}):')
for i, h in enumerate(headers):
    print(f'  [{i}] {h}')

# Count rows
row_count = 0
for _ in ws.iter_rows(min_row=2, values_only=True):
    row_count += 1
print(f'\nData rows: {row_count:,}')

# First 5 rows - need to re-iterate
ws = wb['Public Register']
print('\nFirst 5 rows:')
for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=6, values_only=True)):
    vals = [str(c)[:50] if c else '' for c in row]
    print(f'  Row {row_idx+2}: {vals}')

# Religious keyword search in charity name
ws = wb['Public Register']
religious_keywords = ['CHURCH', 'MOSQUE', 'ISLAMIC', 'CATHOLIC', 'BAPTIST',
                      'METHODIST', 'LUTHERAN', 'EPISCOPAL', 'PRESBYTERIAN',
                      'ORTHODOX', 'PENTECOSTAL', 'EVANGELICAL', 'GOSPEL',
                      'CHRISTIAN', 'SYNAGOGUE', 'TEMPLE', 'HINDU', 'BUDDHIST',
                      'SIKH', 'GURDWARA', 'MASJID', 'CHAPEL', 'CATHEDRAL',
                      'MONASTERY', 'CONVENT', 'PRIORY', 'ABBEY', 'MISSION',
                      'PARISH', 'DIOCESE', 'MINISTRY', 'SALVATION ARMY',
                      'QUAKER', 'UNITARIAN', 'LDS', 'MORMON', 'JEHOVAH',
                      'SHALOM', 'YESHIVA', 'CHRISTADELPHIAN', 'MENNONITE',
                      'CROSS', 'CALVARY', 'BIBLE', 'PRAYER', 'FAITH']

# Find name column
name_idx = None
for i, h in enumerate(headers):
    if h and 'name' in str(h).lower():
        name_idx = i
        break

religious_count = 0
religious_rows = []
total = 0
if name_idx is not None:
    print(f'\nName column: [{name_idx}] {headers[name_idx]}')
    for row in ws.iter_rows(min_row=2, values_only=True):
        total += 1
        if row[name_idx] and any(kw in str(row[name_idx]).upper() for kw in religious_keywords):
            religious_count += 1
            if len(religious_rows) < 20:
                religious_rows.append((str(row[name_idx])[:60], str(row[0])[:20] if row[0] else ''))

print(f'\nTotal rows: {total:,}')
print(f'Religious-sounding names: {religious_count:,}')

print('\nSample religious charities:')
for name, reg in religious_rows:
    print(f'  [{reg}] {name}')

# Also count "CHURCH" specifically
if name_idx is not None:
    ws = wb['Public Register']
    church_count = 0
    mosque_count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        n = str(row[name_idx]).upper() if row[name_idx] else ''
        if 'CHURCH' in n:
            church_count += 1
        if 'MOSQUE' in n or 'ISLAMIC' in n or 'MASJID' in n:
            mosque_count += 1
    print(f'\n"Church" in name: {church_count:,}')
    print(f'"Mosque/Islamic/Masjid" in name: {mosque_count:,}')

wb.close()
